# -*- coding: utf-8 -*-
"""
糖知 SugarGuard AI —— 编码规范文档生成脚本
输出：编码规范.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

OUT_DIR = r"E:\code\Anroid\MyApplication\综合实践\说明"

doc = Document()

_normal = doc.styles["Normal"]
_normal.font.name = "宋体"
_normal.font.size = Pt(11)
_normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
_normal.paragraph_format.line_spacing = 1.5

for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)


def set_run(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def heading(text, level=1, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    sizes = {0: 22, 1: 16, 2: 14, 3: 12}
    set_run(r, "黑体", sizes.get(level, 12), True, (0x1F, 0x3A, 0x5F))
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def para(text, indent=True, size=11, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)


def bullet(text, size=11):
    p = doc.add_paragraph()
    r = p.add_run("• " + text)
    set_run(r, size=size)
    p.paragraph_format.left_indent = Cm(0.5)


def code_block(text, size=9):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, "Consolas", size, color=(0x33, 0x33, 0x33))
    p.paragraph_format.left_indent = Cm(1.0)
    from docx.oxml import OxmlElement
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F5F5F5")
    pPr.append(shd)


def set_cell_shade(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    )


def make_table(headers, rows, col_widths=None, head_color="1F3A5F", font_size=10):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run(r, "黑体", font_size, True, (0xFF, 0xFF, 0xFF))
        set_cell_shade(c, head_color)
    for ri, row in enumerate(rows):
        for j, v in enumerate(row):
            c = table.rows[ri + 1].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(str(v))
            set_run(r, size=font_size)
            if ri % 2 == 1:
                set_cell_shade(c, "F5F7FA")
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


# ==================== 封面 ====================
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(100)
r = cover.add_run("糖知 SugarGuard AI")
set_run(r, "黑体", 26, True, (0x1F, 0x3A, 0x5F))
c2 = doc.add_paragraph()
c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = c2.add_run("编码规范文档")
set_run(r, "黑体", 24, True, (0x26, 0xA6, 0x9A))

for _ in range(4):
    doc.add_paragraph()

info = [
    ("文档名称", "SugarGuard AI 编码规范"),
    ("适用范围", "Android (Kotlin) + 后端 (Java) + AI 服务 (Python)"),
    ("编制人", "鞠星冉"),
    ("编制日期", "2026-04-18"),
    ("版本", "V1.0"),
]
make_table(["项目", "内容"], info, col_widths=[4, 12])

doc.add_page_break()

# ==================== 目录 ====================
heading("目 录", 0, True)
toc = [
    "一、总则",
    "二、Kotlin 编码规范（Android 端）",
    "三、Java 编码规范（Spring Boot 后端）",
    "四、Python 编码规范（AI 服务）",
    "五、通用规范",
    "六、Git Commit 规范",
    "七、代码审查 Checklist",
]
for t in toc:
    p = doc.add_paragraph()
    r = p.add_run(t)
    set_run(r, size=12)
    p.paragraph_format.left_indent = Cm(0.5)

doc.add_page_break()

# ==================== 一、总则 ====================
heading("一、总则")
para("本文档为糖知 SugarGuard AI 项目的统一编码规范，适用于 Android 客户端（Kotlin）、"
     "Spring Boot 后端（Java）和 Python AI 服务三个子项目。所有团队成员在编写代码时应遵循本规范，"
     "代码审查（Code Review）以本规范为评审基准。")

heading("1.1 规范来源", 2)
make_table(
    ["端/层", "规范基线", "工具强制"],
    [
        ["Android (Kotlin)", "Google Kotlin Style Guide", "ktlint + IDE 格式化"],
        ["后端 (Java)", "阿里巴巴 Java 开发手册（黄山版）", "CheckStyle + IDE 格式化"],
        ["AI 服务 (Python)", "PEP 8 + Black 格式化器", "black + flake8 + isort"],
    ],
    col_widths=[3, 6, 5],
)

heading("1.2 约定说明", 2)
bullet("【必须】—— 违反此条将在代码审查中被驳回，必须修改后方可合并")
bullet("【建议】—— 推荐遵循，但可在合理理由下偏离（需在 PR 中说明）")
bullet("【禁止】—— 严禁出现，发现即修改")

# ==================== 二、Kotlin ====================
doc.add_page_break()
heading("二、Kotlin 编码规范（Android 端）")
para("基于 Google Kotlin Style Guide（https://developer.android.com/kotlin/style-guide），"
     "结合项目 Jetpack Compose + MVVM 架构特点进行细化。")

heading("2.1 文件结构", 2)
bullet("【必须】每个 .kt 文件只包含一个顶级类或一组紧密关联的扩展函数")
bullet("【必须】文件名与主要类名一致，使用 PascalCase（如 HomeScreen.kt、LocalMealViewModel.kt）")
bullet("【必须】文件头部按以下顺序排列：package 声明 → 空行 → import 语句 → 空行 → 顶级声明")
bullet("【禁止】使用通配符 import（import xxx.*），IDE 自动优化 import")

heading("2.2 命名规范", 2)
make_table(
    ["元素", "规则", "示例"],
    [
        ["包名", "全小写，不使用下划线", "com.example.myapplication.ui.compose"],
        ["类 / 接口 / 对象", "PascalCase", "HomeScreen, LocalMealViewModel"],
        ["函数", "camelCase", "addMeal(), fetchDailyStats()"],
        ["@Composable 函数", "PascalCase（视为 UI 组件）", "DiaryScreen(), SugarProgressRing()"],
        ["常量（const val / companion）", "UPPER_SNAKE_CASE", "MAX_SUGAR_LIMIT, BASE_URL"],
        ["普通变量 / 属性", "camelCase", "userName, dailySugar"],
        ["LiveData / StateFlow", "_前缀私有 + 公开只读", "_meals / meals"],
        ["布尔型", "is/has/can/should 前缀", "isLoggedIn, hasPermission"],
    ],
    col_widths=[4, 4, 6],
    font_size=9,
)

heading("2.3 格式要求", 2)
bullet("【必须】缩进使用 4 个空格，禁止 Tab")
bullet("【必须】行宽 ≤ 120 字符（IDE 设置 Hard wrap）")
bullet("【必须】大括号 K&R 风格（左括号在行尾，右括号独占一行）")
bullet("【必须】函数参数 > 3 个时，每个参数独占一行并对齐")
bullet("【建议】链式调用每个 . 换行对齐")

heading("2.4 Compose 专项", 2)
bullet("【必须】@Composable 函数用 PascalCase 命名，参数用 modifier: Modifier = Modifier 作为第一个可选参数")
bullet("【必须】State 管理用 remember + mutableStateOf，避免在 Composable 内直接调用 ViewModel 副作用")
bullet("【建议】复杂 Composable 拆分为子组件，单个函数体 ≤ 80 行")
bullet("【禁止】在 LaunchedEffect 中执行无 key 的副作用（必须提供有意义的 key）")

heading("2.5 ViewModel & Repository", 2)
bullet("【必须】ViewModel 中的可变状态使用 private MutableLiveData/MutableStateFlow，对外暴露只读版本")
bullet("【必须】Repository 方法返回 Flow 或 suspend 函数，禁止在 Repository 层切换线程")
bullet("【必须】网络请求异常统一在 ViewModel 层 catch，通过 LiveData/StateFlow 通知 UI")
bullet("【禁止】在 ViewModel 中直接引用 Activity/Context（使用 AndroidViewModel 或 Hilt 注入）")

heading("2.6 Room 数据库", 2)
bullet("【必须】Entity 类使用 @Entity(tableName = \"snake_case\") 注解")
bullet("【必须】DAO 接口方法使用 suspend 或返回 Flow")
bullet("【禁止】在主线程执行数据库操作")

# ==================== 三、Java ====================
doc.add_page_break()
heading("三、Java 编码规范（Spring Boot 后端）")
para("基于《阿里巴巴 Java 开发手册（黄山版）》，结合 Spring Boot 3.2 + JPA + MySQL 技术栈细化。")

heading("3.1 命名规范", 2)
make_table(
    ["元素", "规则", "示例"],
    [
        ["包名", "全小写，点分隔", "com.example.usermanagement.controller"],
        ["类名", "PascalCase，名词", "UserController, MealService"],
        ["接口名", "PascalCase，不加 I 前缀", "UserRepository（非 IUserRepository）"],
        ["方法名", "camelCase，动词开头", "findByUsername(), createMealRecord()"],
        ["常量", "UPPER_SNAKE_CASE", "JWT_SECRET_KEY, TOKEN_EXPIRY"],
        ["局部变量", "camelCase", "mealList, totalSugar"],
        ["DTO 类", "以 Request/Response/DTO 结尾", "LoginRequest, MealRecordDTO"],
        ["枚举值", "UPPER_SNAKE_CASE", "BREAKFAST, LUNCH, DINNER"],
    ],
    col_widths=[3, 5, 6],
    font_size=9,
)

heading("3.2 代码格式", 2)
bullet("【必须】缩进 4 个空格，禁止 Tab")
bullet("【必须】行宽 ≤ 120 字符")
bullet("【必须】左大括号不换行（K&R 风格）")
bullet("【必须】if/else/for/while/do 语句必须使用大括号，即使只有一行")
bullet("【必须】方法间空一行，类内逻辑分组间空一行")

heading("3.3 注释规范", 2)
bullet("【必须】所有 public 类和 public 方法必须有 Javadoc 注释")
bullet("【必须】Controller 的每个端点方法注明：功能描述、@param、@return、HTTP 状态码")
bullet("【禁止】提交注释掉的代码（Dead Code），使用 Git 历史回溯")
bullet("【禁止】使用 // TODO 未标注负责人和预期完成时间")

heading("3.4 Spring Boot 专项", 2)
bullet("【必须】Controller 层仅做参数校验和请求转发，业务逻辑放在 Service 层")
bullet("【必须】Service 层使用 @Transactional 注解管理事务边界")
bullet("【必须】异常使用自定义 BusinessException，统一通过 @ControllerAdvice 处理")
bullet("【必须】REST API 返回统一格式：{code, message, data}")
bullet("【禁止】在 Controller 层直接操作 Repository")
bullet("【禁止】在代码中硬编码数据库密码、API Key（使用 application.yml 或环境变量）")

heading("3.5 JPA & 数据库", 2)
bullet("【必须】Entity 使用 @Table(name = \"snake_case\") 注解指定表名")
bullet("【必须】字段名 camelCase，对应列名 snake_case（通过 naming-strategy 配置）")
bullet("【必须】关联关系明确 fetch = FetchType.LAZY，避免 N+1 查询")
bullet("【建议】复杂查询使用 @Query + JPQL 或原生 SQL，避免方法名过长")

heading("3.6 安全规范", 2)
bullet("【必须】密码使用 BCrypt 加密存储，禁止明文")
bullet("【必须】JWT Token 设置合理过期时间（默认 24h）")
bullet("【必须】接口权限通过 Spring Security 配置，非公开接口必须鉴权")
bullet("【禁止】在日志中输出用户密码、Token 等敏感信息")

# ==================== 四、Python ====================
doc.add_page_break()
heading("四、Python 编码规范（AI 服务）")
para("基于 PEP 8 规范 + Black 格式化器（行宽 88），结合 FastAPI + PyTorch + LangChain 技术栈细化。")

heading("4.1 格式化工具链", 2)
make_table(
    ["工具", "用途", "配置"],
    [
        ["black", "代码格式化（行宽 88）", "pyproject.toml: [tool.black] line-length = 88"],
        ["isort", "import 排序", "profile = \"black\"（兼容 black）"],
        ["flake8", "静态检查", "max-line-length = 88, ignore = E203,W503"],
        ["mypy", "类型检查（可选）", "strict = false（渐进式引入）"],
    ],
    col_widths=[2.5, 4, 7.5],
    font_size=9,
)

heading("4.2 命名规范", 2)
make_table(
    ["元素", "规则", "示例"],
    [
        ["模块/文件名", "snake_case", "deepseek_agent.py, rag_knowledge.py"],
        ["类名", "PascalCase", "DeepSeekAgent, DrinkRecognizer"],
        ["函数/方法", "snake_case", "recognize_drink(), get_health_analysis()"],
        ["常量", "UPPER_SNAKE_CASE", "DEEPSEEK_API_KEY, MODEL_PATH"],
        ["变量", "snake_case", "drink_name, sugar_content"],
        ["私有属性", "单下划线前缀", "_db_session, _model"],
        ["类型别名", "PascalCase", "DrinkList = List[DrinkItem]"],
        ["FastAPI 路由函数", "snake_case 动词", "recognize_drink(), chat_with_ai()"],
    ],
    col_widths=[3, 4, 7],
    font_size=9,
)

heading("4.3 代码风格", 2)
bullet("【必须】缩进 4 个空格，禁止 Tab")
bullet("【必须】行宽 ≤ 88 字符（Black 默认）")
bullet("【必须】字符串统一使用双引号 \"\"（Black 默认）")
bullet("【必须】import 按标准库 → 第三方库 → 本地模块顺序排列，组间空行")
bullet("【必须】函数参数和返回值添加类型注解")
bullet("【禁止】使用 from xxx import *")

heading("4.4 类型注解", 2)
para("所有公开函数必须添加参数类型和返回类型注解：")
code_block('def recognize_drink(\n    image: UploadFile,\n    user_id: int,\n    confidence_threshold: float = 0.5,\n) -> dict[str, Any]:\n    """识别饮品并返回营养信息"""')

heading("4.5 FastAPI 专项", 2)
bullet("【必须】路由使用 RESTful 风格：GET 查询、POST 创建、PUT 更新、DELETE 删除")
bullet("【必须】请求/响应模型使用 Pydantic BaseModel，字段添加 Field(description=...)")
bullet("【必须】异常使用 HTTPException 并指定合理的 status_code")
bullet("【必须】异步函数使用 async def，数据库操作使用异步驱动（aiomysql）")
bullet("【建议】依赖注入使用 Depends()，避免全局变量")

heading("4.6 AI / ML 专项", 2)
bullet("【必须】模型文件路径通过 Settings（pydantic-settings）配置，禁止硬编码")
bullet("【必须】GPU/CPU 自动检测，提供 CPU 降级方案")
bullet("【必须】推理结果包含置信度分数，前端展示时供二次确认")
bullet("【建议】大模型 prompt 模板使用常量或单独文件管理，便于调优迭代")
bullet("【禁止】在生产代码中使用 print() 输出调试信息（使用 loguru）")

# ==================== 五、通用规范 ====================
doc.add_page_break()
heading("五、通用规范")

heading("5.1 日志规范", 2)
make_table(
    ["端", "日志框架", "日志级别使用"],
    [
        ["Android", "Timber", "d=调试 / i=业务流程 / w=异常但可恢复 / e=错误"],
        ["Spring Boot", "SLF4J + Logback", "debug=调试 / info=业务 / warn=告警 / error=异常"],
        ["Python AI", "loguru", "debug=调试 / info=业务 / warning=告警 / error=异常"],
    ],
    col_widths=[3, 4, 7],
)
bullet("【禁止】使用 System.out.println / print() / Log.d() 在生产代码中输出")
bullet("【禁止】在日志中输出密码、Token、完整请求体等敏感信息")
bullet("【必须】异常日志必须包含堆栈信息（logger.error(\"msg\", e)）")

heading("5.2 异常处理", 2)
bullet("【禁止】空 catch 块（catch (Exception e) {}），必须至少记录日志")
bullet("【禁止】用 try-catch 控制业务流程（应使用条件判断）")
bullet("【必须】业务异常使用自定义异常类，包含错误码和用户友好消息")

heading("5.3 安全规范", 2)
bullet("【禁止】在代码中硬编码 API Key、数据库密码、JWT Secret 等敏感信息")
bullet("【必须】敏感配置通过环境变量或 .env 文件注入，.env 加入 .gitignore")
bullet("【必须】用户输入在使用前进行校验和转义，防止 SQL 注入和 XSS")

heading("5.4 版本控制", 2)
bullet("【必须】.gitignore 排除：build 产物、IDE 配置、.env、日志文件、__pycache__")
bullet("【禁止】提交编译产物（.class/.apk/.pyc）、大文件（>10MB）、密钥文件")
bullet("【必须】每次提交前运行格式化工具（ktlint/checkstyle/black）")

# ==================== 六、Git Commit ====================
heading("六、Git Commit 规范")
para("采用 Conventional Commits 格式：<type>(<scope>): <subject>")

make_table(
    ["type", "含义", "示例"],
    [
        ["feat", "新功能", "feat(auth): 添加 JWT 登录接口"],
        ["fix", "修复 Bug", "fix(diary): 修复日记重复添加问题"],
        ["docs", "文档变更", "docs: 更新 README 本地运行指南"],
        ["style", "格式调整（不影响逻辑）", "style(kotlin): ktlint 格式化"],
        ["refactor", "重构（不新增功能/修复 Bug）", "refactor(vm): 提取公共 BaseViewModel"],
        ["test", "添加/修改测试", "test(dao): 新增 MealDao 单元测试"],
        ["chore", "构建/CI/依赖变更", "chore: 升级 Spring Boot 到 3.2.5"],
        ["perf", "性能优化", "perf(ai): ViT 推理 batch_size=1 优化"],
    ],
    col_widths=[2, 4, 8],
    font_size=9,
)

bullet("【必须】scope 为模块名：auth / diary / recognition / ai / analysis / profile / notification 等")
bullet("【必须】subject 使用中文或英文均可，简明扼要（≤50 字符）")
bullet("【建议】关联 TAPD 缺陷/需求号：fix(diary): 修复日记重复 [TAPD-BUG022]")

# ==================== 七、代码审查 ====================
heading("七、代码审查 Checklist")
para("每次 PR（Pull Request）合并前，Reviewer 按以下清单逐项检查：")

make_table(
    ["类别", "检查项", "判定"],
    [
        ["命名", "类/方法/变量命名是否符合本规范", "通过/驳回"],
        ["格式", "是否通过格式化工具（ktlint/checkstyle/black）", "通过/驳回"],
        ["逻辑", "是否存在明显逻辑错误或边界未处理", "通过/驳回"],
        ["安全", "是否有硬编码密钥、SQL 注入风险", "通过/驳回"],
        ["日志", "是否使用正确的日志框架和级别", "通过/驳回"],
        ["异常", "异常是否被正确处理（非空 catch、友好提示）", "通过/驳回"],
        ["测试", "关键逻辑是否有对应单元测试", "通过/建议"],
        ["注释", "公开 API 是否有 Javadoc / docstring", "通过/建议"],
        ["性能", "是否有 N+1 查询、内存泄漏等性能隐患", "通过/建议"],
        ["TAPD", "PR 标题是否含 [TAPD-xxx] 关联号", "通过/驳回"],
    ],
    col_widths=[2, 8, 3],
    font_size=9,
)

para("Reviewer 至少 1 人 Approve 方可合并（核心模块如认证、数据库、AI 接口需 ≥ 2 人）。"
     "驳回项必须修复后重新提交审查。", indent=True)

# ==================== 保存 ====================
out_path = os.path.join(OUT_DIR, "编码规范.docx")
tmp_path = os.path.join(OUT_DIR, "编码规范_tmp.docx")
doc.save(tmp_path)
try:
    os.replace(tmp_path, out_path)
    print("OK →", out_path)
except PermissionError:
    print("OK (tmp) →", tmp_path)
